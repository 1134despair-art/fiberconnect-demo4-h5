from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"c:\Users\EDY\Desktop\工作文件\嗨创\入职培训资料\demo\demo4\APP")
OUTPUT_DIR = ROOT / "智能otdr需求汇总"
DOCX_PATH = OUTPUT_DIR / "智能OTDR独有功能开发梳理.docx"
MD_PATH = OUTPUT_DIR / "智能OTDR独有功能开发梳理.md"


EXCLUDED_SCOPE = [
    "光源 `light-source`：之前设备已有发光控制能力，当前智能 OTDR 页面可以在原有能力基础上复用。",
    "红光源 `vfl`：之前设备已有独立 VFL 页面，智能 OTDR 直接沿用即可。",
    "光功率计 `smart-opm`：虽然智能 OTDR 页面做了硬件化收口，但核心测量模型可复用旧设备 OPM 逻辑。",
    "基础通用壳层能力：如页面头部、部分卡片容器、通用按钮样式等可继续沿用既有设备的基础组件风格。",
]


UNIQUE_MODULES = [
    {
        "name": "智能OTDR独立测量工作台",
        "route": "/device/:id/smart-otdr",
        "files": [
            "src/pages/device/operations/SmartOtdrMeasure.vue",
            "src/pages/device/components/SmartOtdrOverviewPanel.vue",
            "src/pages/device/components/SmartOtdrParameterPanel.vue",
            "src/pages/device/components/SmartOtdrThresholdPanel.vue",
            "src/pages/device/components/SmartOtdrResultPanel.vue",
        ],
        "why_unique": "这里不仅是页面壳层不同，而是智能 OTDR 的 OTDR 主流程本身就需要独立开发。旧设备虽然有便携 OTDR 页面，但不能直接拿来复用，因为智能 OTDR 要求把概览、参数设置、门限设置、测试结果、事件分析入口全部收在同一路由内，并且业务流程和交互结构都不同。",
        "current_state": "当前已经完成主页面结构、测试三态、参数草稿与保存值区分、门限设置、结果页和事件分析入口，但文档口径上必须将该模块视为“需要单独开发”的主模块，而不是旧 OTDR 能力复用。",
        "dev_focus": [
            "后续开发必须按智能 OTDR 独立工作台继续推进，不要直接套便携 OTDR 的旧页面和旧交互。",
            "继续完善一体化流程，不再拆回传统工具页。",
            "保持“保存后的参数优先参与测试”的业务规则。",
            "后续如果接设备通信，需要优先对接这一页的状态机、配置保存、测试结果回填和事件分析跳转。",
        ],
    },
    {
        "name": "事件地图与事件分析链路",
        "route": "/device/:id/event-map",
        "files": [
            "src/pages/device/operations/EventMapOperation.vue",
            "src/pages/device/components/SmartOtdrOverviewPanel.vue",
        ],
        "why_unique": "前三类设备里没有“事件地图”这个独立页面，也没有 OTDR 主页跳转到事件地图做分析的完整链路。",
        "current_state": "当前已具备结果汇总条、轨迹主图、完整事件表、设置菜单、测试、保存和返回 OTDR 的流程。",
        "dev_focus": [
            "后续应把事件地图与 OTDR 测试结果做真实联动，而不是停留在独立 Mock。",
            "建议继续补齐事件点高亮、历史结果加载、事件地图导出和曲线定位联动。",
            "这是智能 OTDR 最典型的独有能力，不能直接从旧设备页面复用。",
        ],
    },
    {
        "name": "三合一网线测试",
        "route": "/device/:id/cable-test",
        "files": [
            "src/pages/device/operations/CableTestOperation.vue",
            "src/pages/device/operations/CableWiremapOperation.vue",
            "src/pages/device/operations/CableSeekOperation.vue",
        ],
        "why_unique": "旧设备里原本是“网络对线”和“网络寻线”两个独立入口，且没有统一的“线序 + 线长 + 寻线”三合一菜单。",
        "current_state": "当前已改造成单入口三模式结构，支持线序、线长、寻线三个能力。",
        "dev_focus": [
            "后续开发要以三合一菜单为准，不再沿用旧的双入口结构。",
            "真实设备接入时，需要把线长模式单独接到硬件数据，而不是复用旧 wiremap/seek 的返回结构。",
            "如果要继续扩展频率、远端模块识别等能力，也应放在这个统一页面里做。",
        ],
    },
    {
        "name": "损耗测试",
        "route": "/device/:id/loss-test",
        "files": [
            "src/pages/device/operations/SmartLossTest.vue",
        ],
        "why_unique": "前三类设备中没有独立的损耗测试页面与流程，不能直接拿旧设备页面替代。",
        "current_state": "当前页面已按硬件收口为独立功能，支持开启/关闭、1310/1550 波长切换、设为参考、参考置零和退出。",
        "dev_focus": [
            "后续需要接真实损耗测试流程，包括发光端与接收端联动、结果计算和结果保存。",
            "当前页面更偏功能壳层和 UI 流程，仍需补真实业务闭环。",
            "这是智能 OTDR 新增工具页，应该单独评估接口和数据结构。",
        ],
    },
    {
        "name": "统一文件管理中心",
        "route": "/device/:id/files",
        "files": [
            "src/pages/device/operations/SmartFileManager.vue",
        ],
        "why_unique": "旧设备有记录页，但没有面向智能 OTDR 的统一文件中心来管理 OTDR 结果、事件地图结果、损耗测试结果，并切换内部存储 / SD 卡 / U 盘。",
        "current_state": "当前已具备目录信息卡、分组、工具栏、文件列表和多存储切换 UI。",
        "dev_focus": [
            "后续需要打通真实文件介质、文件加载、重命名、移动、删除和导出。",
            "建议把 OTDR、事件地图、损耗测试统一纳入这一页管理，而不是各页各自维护一套历史列表。",
            "这是智能 OTDR 的平台化能力，属于旧设备中没有的统一管理域。",
        ],
    },
    {
        "name": "智能OTDR系统设置扩展项",
        "route": "/device/:id/system",
        "files": [
            "src/pages/device/operations/SystemSettings.vue",
        ],
        "why_unique": "虽然旧设备也有系统设置页，但智能 OTDR 在现有项目里已经扩展出一组需要明确梳理的专项能力，不能简单归类为“可忽略的复用项”。",
        "current_state": "当前系统设置中已覆盖屏幕亮度、触屏提示音、蓝牙开关、休眠、休眠时间、自动关机、关机时间、语言、日期时间、系统信息、固件升级、恢复出厂，以及智能 OTDR 专属的 4G 联网、物联网卡号、电池电量、手电筒、截屏功能、远程升级等内容。",
        "dev_focus": [
            "后续要把智能 OTDR 系统设置作为独立开发范围明确给同事，尤其是蓝牙、休眠/自动关机、语言/时间、4G、物联网卡、电池、手电筒、截屏、远程升级、背光校正等功能。",
            "如果接真实设备，这一页要负责系统级设置下发、状态读取、升级状态反馈和本地能力联动。",
            "文档和排期上不应再把智能 OTDR 系统设置整体忽略掉。",
        ],
    },
    {
        "name": "设备激活与授权管理",
        "route": "/device/:id 详情页 + /device/:id/system",
        "files": [
            "src/pages/device/components/SmartOtdrActivationCard.vue",
            "src/pages/device/DeviceDetail.vue",
            "src/pages/device/operations/SystemSettings.vue",
        ],
        "why_unique": "智能 OTDR 的设备激活不是普通状态展示，而是主详情页激活卡、蓝牙依赖校验、激活码录入、扫码填充、授权到期展示、授权时间设置、系统设置授权入口共同组成的一套独立业务流程。",
        "current_state": "当前已具备详情页激活卡、状态徽标、授权到期时间、最近激活时间、激活码输入、扫码模拟、清空、蓝牙校验、激活成功回写，以及系统设置页中的设备激活与授权设置入口。",
        "dev_focus": [
            "后续要把设备激活单独列入开发范围，不能遗漏。",
            "需要继续补齐真实激活接口、授权校验、到期锁机策略、授权时间配置和系统设置联动。",
            "激活结果会影响智能 OTDR 详情页展示与使用状态，属于核心业务模块。",
        ],
    },
]


NEXT_ACTIONS = [
    "后续开发同事应优先从这些独有模块入手，不要把精力花在已经能复用的光源、红光源、光功率计等能力上。",
    "OTDR 主流程在智能 OTDR 上应被视为独立开发模块，不能按旧便携 OTDR 页面直接复用。",
    "如果要接设备通信，优先定义智能 OTDR 主流程、事件地图、损耗测试、统一文件管理、系统设置、设备激活的接口协议与状态模型。",
    "智能 OTDR 主详情页和事件地图之间已经形成联动关系，后续开发需要把两页视为同一业务链路处理。",
    "三合一网线测试要坚持单入口方案，不能再回退到旧的两个入口结构。",
    "系统设置和设备激活这两块在交接文档中必须显式列出，避免后续排期遗漏。",
]


SOURCE_FILES = [
    "src/pages/device/operations/SmartOtdrMeasure.vue",
    "src/pages/device/components/SmartOtdrOverviewPanel.vue",
    "src/pages/device/components/SmartOtdrParameterPanel.vue",
    "src/pages/device/components/SmartOtdrThresholdPanel.vue",
    "src/pages/device/components/SmartOtdrResultPanel.vue",
    "src/pages/device/operations/EventMapOperation.vue",
    "src/pages/device/operations/CableTestOperation.vue",
    "src/pages/device/operations/CableWiremapOperation.vue",
    "src/pages/device/operations/CableSeekOperation.vue",
    "src/pages/device/operations/SmartLossTest.vue",
    "src/pages/device/operations/SmartEndfaceInspect.vue",
    "src/pages/device/operations/SmartFileManager.vue",
    "src/pages/device/operations/SystemSettings.vue",
    "src/pages/device/components/SmartOtdrActivationCard.vue",
    "src/pages/device/DeviceDetail.vue",
    "src/router/index.ts",
]


def set_run_font(run, bold=False, size=10.5):
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    set_run_font(run, bold=bold)
    p.paragraph_format.space_after = Pt(0)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        shade_cell(table.rows[0].cells[i], "DCEBFF")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    return table


def build_markdown():
    lines = []
    lines.append("# 智能OTDR独有功能开发梳理")
    lines.append("")
    lines.append("- 文档目标：只梳理智能 OTDR 中无法直接复用前三类设备现有页面的独有功能。")
    lines.append("- 前三类设备范围：便携 OTDR、多合一光功率计、熔接机。")
    lines.append("")
    lines.append("## 已排除的可复用能力")
    for item in EXCLUDED_SCOPE:
        lines.append(f"- {item}")
    lines.append("")
    lines.append("## 需要单独开发或持续深化的智能OTDR独有模块")
    lines.append("| 模块 | 路由 | 独有原因 | 当前状态 | 后续开发重点 |")
    lines.append("| --- | --- | --- | --- | --- |")
    for item in UNIQUE_MODULES:
        lines.append(
            f"| {item['name']} | `{item['route']}` | {item['why_unique']} | {item['current_state']} | {'；'.join(item['dev_focus'])} |"
        )
    lines.append("")
    lines.append("## 建议的开发优先级")
    for item in NEXT_ACTIONS:
        lines.append(f"- {item}")
    lines.append("")
    lines.append("## 代码依据")
    for item in SOURCE_FILES:
        lines.append(f"- `{item}`")
    lines.append("")
    return "\n".join(lines)


def build_docx():
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("智能OTDR独有功能开发梳理")
    set_run_font(run, bold=True, size=18)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("仅保留智能 OTDR 相对前三类设备新增的功能域")
    set_run_font(run, size=10.5)

    p = document.add_paragraph()
    run1 = p.add_run("前三类设备范围：")
    set_run_font(run1, bold=True)
    run2 = p.add_run("便携 OTDR、多合一光功率计、熔接机。")
    set_run_font(run2)

    p2 = document.add_paragraph()
    run1 = p2.add_run("本次排除原则：")
    set_run_font(run1, bold=True)
    run2 = p2.add_run("凡是旧设备里已有成熟页面或成熟能力模型可直接复用的模块，本报告不再展开。")
    set_run_font(run2)

    document.add_heading("1. 已排除的可复用能力", level=1)
    add_bullets(document, EXCLUDED_SCOPE)

    document.add_heading("2. 需要单独开发或持续深化的智能OTDR独有模块", level=1)
    rows = []
    for item in UNIQUE_MODULES:
        rows.append((
            item["name"],
            item["route"],
            item["why_unique"],
            item["current_state"],
            "；".join(item["dev_focus"]),
        ))
    add_table(document, ["模块", "路由", "独有原因", "当前状态", "后续开发重点"], rows)

    document.add_heading("3. 推荐开发顺序", level=1)
    add_bullets(document, NEXT_ACTIONS)

    document.add_heading("4. 结论", level=1)
    conclusions = [
        "后续智能 OTDR 开发不应再从光源、红光源、光功率计这些已存在能力重新起页，而应优先做旧设备没有的增量模块。",
        "当前真正需要独立投入的重点，是智能 OTDR 独立测量工作台、事件地图、三合一网线测试、损耗测试、统一文件管理、系统设置扩展项、设备激活与授权管理。",
        "其中 OTDR 主流程、事件地图、系统设置扩展项、设备激活是这次用户明确要求补充和改口径的重点，后续排期时必须单独立项。",
    ]
    add_bullets(document, conclusions)

    document.add_heading("5. 代码依据", level=1)
    add_bullets(document, SOURCE_FILES)

    document.save(str(DOCX_PATH))


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    MD_PATH.write_text(build_markdown(), encoding="utf-8")
    build_docx()
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
