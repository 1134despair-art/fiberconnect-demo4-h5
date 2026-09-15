from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"c:\Users\EDY\Desktop\工作文件\嗨创\入职培训资料\demo\demo4\APP")
OUTPUT_DIR = ROOT / "智能otdr需求汇总"
DOCX_PATH = OUTPUT_DIR / "智能OTDR交接书未覆盖功能梳理.docx"
MD_PATH = OUTPUT_DIR / "智能OTDR交接书未覆盖功能梳理.md"
SOURCE_DOC = OUTPUT_DIR / "智能OTDR项目交接书.docx"


ANALYSIS_PRINCIPLES = [
    "本报告只列出“原交接书中未明确提及，但当前系统前端已经具备”的功能、交互或配置能力。",
    "若原交接书已经提到某项能力，但当前系统只是把它做得更细，本报告会归类为“实现细化项”，避免与真正新增功能混淆。",
    "本报告基于当前前端代码与 Mock 交互，不代表底层硬件链路已经全部真实接通。",
]


EXTRA_FEATURE_ROWS = [
    (
        "新增交互入口",
        "全局快捷入口 + 8宫格快捷面板",
        "交接书分别提到手电筒、截屏、蓝牙、4G 等单项能力，但没有描述一个覆盖所有智能 OTDR 页面、可直接调起这些动作的全局快捷入口。",
        "当前系统已提供右下角悬浮快捷按钮，支持自动/手动测试、主页、文件管理、手电筒、截屏、4G、蓝牙、关机 8 项快捷操作。",
        "src/pages/device/components/SmartOtdrQuickPanel.vue；src/pages/device/components/SmartOtdrPageShell.vue",
    ),
    (
        "新增交互增强",
        "快捷入口支持拖拽和位置记忆",
        "交接书未提及快捷入口可拖拽、拖拽边界控制、刷新后位置记忆等移动端交互细节。",
        "当前系统已支持快捷按钮在设备固定宽度主体内拖拽，并通过 localStorage 记住位置。",
        "src/pages/device/components/SmartOtdrQuickPanel.vue",
    ),
    (
        "新增平台能力",
        "文件设置中心与高级命名配置",
        "交接书只提到文件名称格式编辑、文件删除/加载/移动/重命名等常规管理，没有展开独立的“文件设置中心”。",
        "当前系统已提供文件设置页，支持存储路径、按日期创建文件夹、自动保存、存储格式、文件命名方式，以及任务号、光缆编号、起点、终点、方向、注释等高级字段。",
        "src/pages/device/operations/SmartFileManager.vue",
    ),
    (
        "新增文件联动",
        "事件地图文件可直接加载到 OTDR 页面",
        "交接书提到事件地图结果可保存和历史加载，但没有描述“事件地图文件统一加载回 OTDR 测量页”的跨页面联动方式。",
        "当前系统文件管理中 OTDR 文件和事件地图文件都支持点击加载，并统一跳转到 OTDR 页面显示识别结果。",
        "src/pages/device/operations/SmartFileManager.vue；src/pages/device/operations/SmartOtdrMeasure.vue",
    ),
    (
        "新增业务流程增强",
        "设备激活完整操作链路",
        "交接书只写了激活码、蓝牙激活、授权时间编辑、到期锁机，没有展开具体前端操作闭环。",
        "当前系统已具备详情页激活卡、授权有效期、最近激活时间、激活码输入、扫码填充、清空、蓝牙依赖校验、激活成功反馈，以及系统设置里的激活信息与激活入口。",
        "src/pages/device/components/SmartOtdrActivationCard.vue；src/pages/device/DeviceDetail.vue；src/pages/device/operations/SystemSettings.vue",
    ),
    (
        "新增系统项",
        "休眠与休眠时间设置",
        "交接书提到自动关机、背光、蜂鸣、多语言等系统项，但没有提到“休眠”和“休眠时间”这组独立设置。",
        "当前系统系统设置页已提供休眠开关和休眠时间调节。",
        "src/pages/device/operations/SystemSettings.vue",
    ),
    (
        "新增系统项",
        "蓝牙设备管理与激活信息页签",
        "交接书只提到支持蓝牙功能、蓝牙激活、蓝牙升级，没有描述蓝牙设备管理弹层和系统内激活信息页签。",
        "当前系统系统设置内已包含蓝牙连接/删除交互，以及设备序列号、激活状态、激活时间等激活信息展示。",
        "src/pages/device/operations/SystemSettings.vue",
    ),
    (
        "新增分析视图",
        "OTDR 结果“事件对比”视图",
        "交接书提到事件判别及分析、事件地图功能，但没有明确写“事件对比”这一独立结果视图。",
        "当前系统 OTDR 结果区除结果列表、事件列表外，还提供事件对比视图。",
        "src/pages/device/operations/SmartOtdrMeasure.vue",
    ),
    (
        "实现细化项",
        "OTDR 参数选项细化到更高颗粒度",
        "交接书写的是能力范围，例如双波长、100m~200km、实时/平均模式、折射率、门限设置等，没有列出具体所有选项值。",
        "当前系统已细化出 1310+1550 组合波长、最大 260km 量程、15s/30s/实时/自定义测量时间，以及 1310/1550 分波长平均损耗、双散射系数字段。",
        "src/pages/device/components/SmartOtdrParameterPanel.vue；src/pages/device/components/SmartOtdrThresholdPanel.vue；src/pages/device/components/smartOtdrTypes.ts",
    ),
    (
        "实现细化项",
        "网线测试收口为单入口三模式",
        "交接书分别提到线序、线长、寻线能力，但没有规定一定要做成单入口三合一菜单。",
        "当前系统已将网线测试重构为统一入口页，收口线序、线长、寻线三种模式。",
        "src/pages/device/operations/CableTestOperation.vue；src/pages/device/operations/CableWiremapOperation.vue；src/pages/device/operations/CableSeekOperation.vue",
    ),
]


SUMMARY_POINTS = [
    "原交接书已经覆盖了智能 OTDR 的大部分主业务能力，但当前系统在“平台化交互”和“前端配置中心”上已经走得更细。",
    "本次最值得补回交接文档的内容，优先建议加入：全局快捷入口、文件设置中心、设备激活完整链路、休眠设置、OTDR 事件对比视图。",
    "像 260km 量程、1310+1550 组合波长、15s/30s/自定义时间等，应作为“实现细化项”补充到需求附录，而不是单独写成新模块。",
]


SOURCE_FILES = [
    "智能otdr需求汇总/智能OTDR项目交接书.docx",
    "src/pages/device/components/SmartOtdrQuickPanel.vue",
    "src/pages/device/components/SmartOtdrPageShell.vue",
    "src/pages/device/operations/SmartFileManager.vue",
    "src/pages/device/operations/SmartOtdrMeasure.vue",
    "src/pages/device/components/SmartOtdrParameterPanel.vue",
    "src/pages/device/components/SmartOtdrThresholdPanel.vue",
    "src/pages/device/components/smartOtdrTypes.ts",
    "src/pages/device/components/SmartOtdrActivationCard.vue",
    "src/pages/device/DeviceDetail.vue",
    "src/pages/device/operations/SystemSettings.vue",
    "src/pages/device/operations/CableTestOperation.vue",
    "src/pages/device/operations/CableWiremapOperation.vue",
    "src/pages/device/operations/CableSeekOperation.vue",
]


def set_run_font(run, bold=False, size=10.5):
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    set_run_font(run, bold=bold)
    paragraph.paragraph_format.space_after = Pt(0)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        shade_cell(table.rows[0].cells[i], "DCEBFF")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    return table


def build_markdown():
    lines = [
        "# 智能OTDR交接书未覆盖功能梳理",
        "",
        f"- 原始文档：`{SOURCE_DOC.name}`",
        "- 输出目标：梳理原交接书未明确提及，但当前系统已经具备的功能、交互和实现细化项。",
        "",
        "## 判断原则",
    ]
    for item in ANALYSIS_PRINCIPLES:
        lines.append(f"- {item}")
    lines.extend([
        "",
        "## 文档未覆盖但系统已有的内容",
        "| 分类 | 系统现有功能 | 文档缺口说明 | 当前系统表现 | 代码依据 |",
        "| --- | --- | --- | --- | --- |",
    ])
    for row in EXTRA_FEATURE_ROWS:
        lines.append(f"| {row[0]} | {row[1]} | {row[2]} | {row[3]} | {row[4]} |")
    lines.extend([
        "",
        "## 结论与建议",
    ])
    for item in SUMMARY_POINTS:
        lines.append(f"- {item}")
    lines.extend([
        "",
        "## 代码依据",
    ])
    for item in SOURCE_FILES:
        lines.append(f"- `{item}`")
    lines.append("")
    return "\n".join(lines)


def build_docx():
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("智能OTDR交接书未覆盖功能梳理")
    set_run_font(run, bold=True, size=18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("基于原交接书与当前系统代码实现的补充对比报告")
    set_run_font(run, size=10.5)

    doc_info = document.add_paragraph()
    run = doc_info.add_run("原始文档：")
    set_run_font(run, bold=True)
    run = doc_info.add_run(SOURCE_DOC.name)
    set_run_font(run)

    document.add_heading("1. 判断原则", level=1)
    add_bullets(document, ANALYSIS_PRINCIPLES)

    document.add_heading("2. 文档未覆盖但系统已有的内容", level=1)
    add_table(
        document,
        ["分类", "系统现有功能", "文档缺口说明", "当前系统表现", "代码依据"],
        EXTRA_FEATURE_ROWS,
    )

    document.add_heading("3. 结论与建议", level=1)
    add_bullets(document, SUMMARY_POINTS)

    document.add_heading("4. 代码依据", level=1)
    add_bullets(document, SOURCE_FILES)

    document.save(str(DOCX_PATH))


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    MD_PATH.write_text(build_markdown(), encoding="utf-8")
    build_docx()
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
